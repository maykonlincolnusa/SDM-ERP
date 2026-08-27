from __future__ import annotations

import hashlib
import mimetypes
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import Settings
from app.firebase import get_firebase_app
from app.repositories.store import Store
from app.schemas import MediaUpdate

SUPPORTED_EXTENSIONS = {
    ".pdf", ".xlsx", ".xls", ".csv", ".doc", ".docx", ".txt",
    ".jpg", ".jpeg", ".png", ".webp",
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".csv"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024


class LocalMediaStorage:
    mode = "local"

    def __init__(self, root: Path):
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)

    def put_file(self, source: Path, key: str) -> None:
        target = self.root / key
        target.parent.mkdir(parents=True, exist_ok=True)
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)

    def put_bytes(self, content: bytes, key: str) -> None:
        target = self.root / key
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)

    def path(self, key: str) -> Path:
        return self.root / key

    def read(self, key: str) -> bytes:
        return self.path(key).read_bytes()

    def delete(self, key: str) -> None:
        self.path(key).unlink(missing_ok=True)


class FirebaseMediaStorage:
    mode = "firebase"

    def __init__(self, settings: Settings):
        try:
            from firebase_admin import storage
        except ImportError as error:  # pragma: no cover - setup-specific
            raise RuntimeError("firebase-admin is required for Firebase Storage") from error
        self.bucket = storage.bucket(settings.firebase_storage_bucket, app=get_firebase_app(settings))

    def put_file(self, source: Path, key: str) -> None:
        blob = self.bucket.blob(key)
        blob.upload_from_filename(str(source), content_type=_mime_type(source.name))

    def put_bytes(self, content: bytes, key: str) -> None:
        self.bucket.blob(key).upload_from_string(content, content_type=_mime_type(key))

    def read(self, key: str) -> bytes:
        return self.bucket.blob(key).download_as_bytes()

    def delete(self, key: str) -> None:  # pragma: no cover - depende do bucket
        blob = self.bucket.blob(key)
        if blob.exists():
            blob.delete()


class MediaService:
    """Catalogs operational files and keeps a stable copy in local or Firebase storage."""

    def __init__(self, store: Store, settings: Settings):
        self.store = store
        self.settings = settings
        self.workspace = Path.cwd().resolve()
        if store.mode == "firebase":
            if not settings.firebase_storage_bucket:
                raise RuntimeError("FIREBASE_STORAGE_BUCKET is required in Firebase mode")
            self.storage: LocalMediaStorage | FirebaseMediaStorage = FirebaseMediaStorage(settings)
        else:
            self.storage = LocalMediaStorage(settings.media_library_path)

    def list(self, category: str | None = None, kind: str | None = None) -> list[dict[str, Any]]:
        records = self.store.list("media")
        if category:
            records = [record for record in records if record.get("category") == category]
        if kind:
            records = [record for record in records if record.get("kind") == kind]
        return sorted(records, key=lambda record: record.get("updated_at", ""), reverse=True)

    def summary(self) -> dict[str, Any]:
        records = self.list()
        categories: dict[str, int] = {}
        for record in records:
            category = record.get("category", "Sem categoria")
            categories[category] = categories.get(category, 0) + 1
        return {
            "total": len(records),
            "images": sum(record.get("kind") == "image" for record in records),
            "needs_review": sum(record.get("status") == "Revisar" for record in records),
            "categories": [{"name": name, "count": count} for name, count in sorted(categories.items())],
            "storage": self.storage.mode,
        }

    def sync_workspace(self) -> dict[str, int]:
        self.settings.data_raw_path.mkdir(parents=True, exist_ok=True)
        self.settings.media_inbox_path.mkdir(parents=True, exist_ok=True)
        added = updated = skipped = 0
        dispensados = {row["id"] for row in self.store.list("media_dismissed")}
        for path, source in self._source_files():
            # Arquivo dispensado pelo gestor nao volta na sincronizacao.
            if _file_hash(path) in dispensados:
                skipped += 1
                continue
            outcome = self._ingest_file(path, source)
            if outcome == "added":
                added += 1
            elif outcome == "updated":
                updated += 1
            else:
                skipped += 1
        return {"added": added, "updated": updated, "skipped": skipped, "total": added + updated + skipped}

    def upload(self, filename: str, content: bytes, category: str | None = None) -> dict[str, Any]:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError("Formato não suportado. Envie PDF, planilha, documento ou imagem.")
        if not content:
            raise ValueError("O arquivo enviado está vazio.")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("Cada arquivo pode ter no máximo 25 MB.")
        checksum = hashlib.sha256(content).hexdigest()
        existing = self._by_checksum().get(checksum)
        if existing:
            existing["updated_at"] = _now()
            self.store.put("media", existing["id"], existing)
            return existing
        record = self._new_record(filename, suffix, checksum, len(content), "upload")
        if category:
            record["category"] = category.strip()[:60]
            record["status"] = "Organizado"
        self.storage.put_bytes(content, record["storage_key"])
        self.store.put("media", record["id"], record)
        return record

    def update(self, media_id: str, payload: MediaUpdate) -> dict[str, Any]:
        record = self.store.get("media", media_id)
        if not record:
            raise ValueError("Arquivo não encontrado.")
        values = payload.model_dump(exclude_unset=True)
        if "tags" in values and values["tags"] is not None:
            values["tags"] = list(dict.fromkeys(tag.strip().lower()[:30] for tag in values["tags"] if tag.strip()))
        if values.get("entity_type") and not values.get("entity_id"):
            raise ValueError("Informe o identificador do vínculo ou remova o tipo de vínculo.")
        if values.get("entity_id") and not values.get("entity_type"):
            raise ValueError("Informe o tipo do vínculo para salvar o identificador.")
        record.update(values)
        record["status"] = "Organizado" if record.get("category") else "Revisar"
        record["updated_at"] = _now()
        self.store.put("media", media_id, record)
        return record

    def delete(self, media_id: str) -> dict[str, str]:
        """Tira o arquivo do acervo e apaga a copia do sistema.

        O original do usuario em `data/raw` nunca e tocado: o ERP catalogou o
        arquivo, nao virou dono dele. Mas o checksum entra numa lista de
        dispensados, senao a proxima sincronizacao traria tudo de volta e a
        exclusao pareceria nao ter funcionado.
        """

        record = self.store.get("media", media_id)
        if not record:
            raise ValueError("Arquivo não encontrado.")
        try:
            self.storage.delete(record["storage_key"])
        except OSError as error:  # pragma: no cover - falha de disco
            raise ValueError(f"Não foi possível remover o arquivo armazenado: {error}") from error
        self.store.delete("media", media_id)
        if record.get("source") != "upload":
            self.store.put("media_dismissed", record["checksum"], {
                "id": record["checksum"],
                "original_name": record.get("original_name"),
                "dismissed_at": _now(),
            })
        return {
            "id": media_id,
            "status": "excluido",
            "detail": (
                "Arquivo removido do acervo. O original na pasta de origem não foi apagado, "
                "mas ele não volta na próxima sincronização."
                if record.get("source") != "upload" else "Arquivo removido do acervo."
            ),
        }

    def get(self, media_id: str) -> dict[str, Any] | None:
        return self.store.get("media", media_id)

    def local_path(self, record: dict[str, Any]) -> Path | None:
        if isinstance(self.storage, LocalMediaStorage):
            path = self.storage.path(record["storage_key"])
            return path if path.exists() else None
        return None

    def content(self, record: dict[str, Any]) -> bytes:
        return self.storage.read(record["storage_key"])

    def _source_files(self) -> list[tuple[Path, str]]:
        candidates: list[tuple[Path, str]] = []
        raw = self.settings.data_raw_path.resolve()
        inbox = self.settings.media_inbox_path.resolve()
        if raw.exists():
            for path in raw.rglob("*"):
                if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                    source = "inbox" if path.is_relative_to(inbox) else "raw"
                    candidates.append((path, source))
        if self.settings.media_sync_legacy_files:
            for path in self.workspace.iterdir():
                if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                    candidates.append((path, "workspace"))
        return sorted(candidates, key=lambda item: item[0].name.lower())

    def _ingest_file(self, path: Path, source: str) -> str:
        checksum = _file_hash(path)
        existing = self._by_checksum().get(checksum)
        if existing:
            existing["last_synced_at"] = _now()
            existing["updated_at"] = _now()
            self.store.put("media", existing["id"], existing)
            return "skipped"
        record = self._new_record(path.name, path.suffix.lower(), checksum, path.stat().st_size, source)
        record["metadata"].update(self._inspect(path, record["kind"]))
        self.storage.put_file(path, record["storage_key"])
        self.store.put("media", record["id"], record)
        return "added"

    def _new_record(self, filename: str, suffix: str, checksum: str, size: int, source: str) -> dict[str, Any]:
        kind, category, status, tags = _classify(filename, suffix)
        created = _now()
        return {
            "id": f"media-{checksum[:20]}",
            "title": Path(filename).stem[:120],
            "original_name": filename,
            "kind": kind,
            "category": category,
            "status": status,
            "tags": tags,
            "entity_type": None,
            "entity_id": None,
            "mime_type": _mime_type(filename),
            "size": size,
            "checksum": checksum,
            "source": source,
            "storage_key": f"media/{created[:7].replace('-', '/')}/{checksum}{suffix}",
            "metadata": _known_metadata(filename),
            "created_at": created,
            "updated_at": created,
            "last_synced_at": created,
        }

    def _by_checksum(self) -> dict[str, dict[str, Any]]:
        return {record.get("checksum"): record for record in self.store.list("media") if record.get("checksum")}

    @staticmethod
    def _inspect(path: Path, kind: str) -> dict[str, Any]:
        if kind != "spreadsheet" or path.suffix.lower() not in {".xlsx", ".xls"}:
            return {}
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return {"sheets": workbook.sheetnames[:16], "sheet_count": len(workbook.sheetnames)}
            finally:
                workbook.close()
        except Exception:
            return {"inspection": "Não foi possível ler a estrutura da planilha."}


def _classify(filename: str, suffix: str) -> tuple[str, str, str, list[str]]:
    lowered = filename.lower()
    if "contrato de loca" in lowered:
        return "document", "Contratos", "Revisar", ["contrato", "locação", "validade"]
    if suffix in IMAGE_EXTENSIONS:
        if "whatsapp image" in lowered:
            return "image", "Compras", "Revisar", ["comprovante", "compra", "revisar"]
        return "image", "Imagens", "Revisar", ["imagem", "revisar"]
    if suffix in SPREADSHEET_EXTENSIONS:
        if "precifica" in lowered:
            return "spreadsheet", "Precificação", "Organizado", ["preço", "ficha-técnica"]
        if "finance" in lowered:
            return "spreadsheet", "Financeiro", "Organizado", ["financeiro", "histórico"]
        if "compar" in lowered:
            return "spreadsheet", "Compras", "Revisar", ["comparativo", "fornecedor"]
        return "spreadsheet", "Planilhas", "Revisar", ["planilha", "revisar"]
    return "document", "Documentos", "Revisar", ["documento", "revisar"]


def _known_metadata(filename: str) -> dict[str, Any]:
    """Known records are reviewable hints; they never create financial entries by themselves."""

    lowered = filename.lower()
    if "contrato de loca" in lowered:
        return {
            "document_type": "lease_contract",
            "business_date": "2024-03-20",
            "valid_until": "2025-03-20",
            "monthly_amount": 1000.0,
            "attention": "Vencimento identificado; confirme renovação antes de usar como contrato vigente.",
        }
    receipt_data = {
        "14.10.31": {"document_type": "expense_note", "business_date": "2026-08-05", "attention": "Anotação de despesas a revisar antes de gerar recorrências."},
        "14.32.36": {"document_type": "purchase_receipt", "business_date": "2026-07-15", "supplier": "Nechio Comércio", "total": 1351.35, "attention": "Compra de insumos; confira itens e unidades antes de lançar estoque."},
        "14.33.51": {"document_type": "purchase_receipt", "business_date": "2026-07-15", "supplier": "JME Rauta Comercial", "total": 194.0, "attention": "Compra de complementos; confira itens e unidades antes de lançar estoque."},
        "14.35.00": {"document_type": "purchase_receipt", "business_date": "2026-07-15", "total": 264.29, "attention": "Compra de embalagens; confirme fornecedor antes de lançar estoque."},
        "08.29.15": {"document_type": "purchase_receipt", "business_date": "2026-08-06", "supplier": "Nechio Comércio", "total": 1810.67, "attention": "Compra de insumos; confira itens e unidades antes de lançar estoque."},
        "08.29.35": {"document_type": "purchase_receipt", "business_date": "2026-08-06", "supplier": "Trembão Atacado", "total": 291.43, "attention": "Compra de bebidas e leite; confira itens e unidades antes de lançar estoque."},
        "08.51.29": {"document_type": "purchase_receipt", "business_date": "2026-08-06", "total": 276.74, "attention": "Compra de embalagens; confirme fornecedor antes de lançar estoque."},
    }
    for fragment, metadata in receipt_data.items():
        if fragment in lowered:
            return metadata
    return {}


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _mime_type(filename: str) -> str:
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()
