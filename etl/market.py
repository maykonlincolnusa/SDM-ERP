"""Competitive market workbook transformation with auditable normalization."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .transform import as_number, normalized_text


ALIASES = {
    "colibr roxo": "colibri roxo",
    "lha pura": "ilha pura",
    "maca": "maçã",
    "pessego": "pêssego",
}

PRODUCT_REVIEW_ALIASES = {
    "colibr roxo": "colibri roxo",
}


def normalize_market_workbook(
    path: Path,
    source_hash: str,
    source_path: str,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    records: list[dict[str, Any]] = []
    diagnostics: list[dict[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            header_index, columns = _find_market_header(rows)
            if header_index is None:
                diagnostics.append({"sheet": sheet.title, "reason": "market_header_not_found"})
                continue
            for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                locality = _value(row, columns["localidade"])
                product = _value(row, columns["produto"])
                placement = as_number(_value(row, columns["colocacao"]))
                if not locality or not product or not placement or placement <= 0:
                    continue
                base_raw = _value(row, columns["base"])
                fruits_raw = _value(row, columns["frutas"])
                social_raw = _value(row, columns["padrao_social"])
                bases, base_flags = _tokens(base_raw)
                fruits, fruit_flags = _tokens(fruits_raw)
                social = _social_class(social_raw)
                product_normalized = normalized_text(product)
                product_review = PRODUCT_REVIEW_ALIASES.get(product_normalized)
                records.append(
                    {
                        "record_type": "market_observation",
                        "source_hash": source_hash,
                        "source_path": source_path,
                        "source_sheet": sheet.title,
                        "source_row": row_number,
                        "locality_raw": locality,
                        "locality_normalized": normalized_text(locality),
                        "product_raw": product,
                        "product_normalized": product_normalized,
                        "product_review_suggestion": product_review,
                        "base_raw": base_raw,
                        "bases": bases,
                        "fruits_raw": fruits_raw,
                        "fruits": fruits,
                        "social_class_raw": social_raw,
                        "social_class": social,
                        "placement": int(placement),
                        "quality_flags": [
                            *base_flags,
                            *fruit_flags,
                            *([] if social else ["unknown_social_class"]),
                            *([] if not product_review else [f"alias_review:{product}->{product_review}" ]),
                        ],
                    }
                )
    finally:
        workbook.close()
    return records, diagnostics


def _find_market_header(rows: list[tuple[Any, ...]]) -> tuple[int | None, dict[str, int]]:
    targets = {"localidade", "produto", "base", "frutas", "padrao_social", "colocacao"}
    for index, row in enumerate(rows[:10]):
        labels = [normalized_text(value).replace(" ", "_") for value in row]
        columns = {target: next((column for column, label in enumerate(labels) if target in label), None) for target in targets}
        if all(column is not None for column in columns.values()):
            return index, {key: int(value) for key, value in columns.items()}
    return None, {}


def _value(row: tuple[Any, ...], column: int) -> str:
    return str(row[column]).strip() if column < len(row) and row[column] is not None else ""


def _tokens(value: str) -> tuple[list[dict[str, str]], list[str]]:
    entries: list[dict[str, str]] = []
    flags: list[str] = []
    for raw in (part.strip() for part in value.split(";")):
        if not raw:
            continue
        alias_key = normalized_text(raw)
        canonical = ALIASES.get(alias_key, raw.strip())
        entries.append({"raw": raw, "normalized": normalized_text(canonical), "canonical": canonical})
        if canonical != raw:
            flags.append(f"alias_review:{raw}->{canonical}")
    return entries, flags


def _social_class(value: str) -> str | None:
    normalized = normalized_text(value)
    if normalized.startswith("alta"):
        return "Alta"
    if normalized.startswith("media"):
        return "Média"
    if normalized.startswith("baixa"):
        return "Baixa"
    return None
