"""Business transformations from pricing workbooks to silver JSONL records."""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TRANSFORMATION_VERSION = 5


def normalized_text(value: Any) -> str:
    text = str(value or "").strip()
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").lower()


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", normalized_text(value)).strip("-") or "item"


def as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    candidate = value.strip().replace("R$", "").replace(" ", "")
    if not candidate:
        return None
    if "," in candidate:
        candidate = candidate.replace(".", "").replace(",", ".")
    try:
        return float(candidate)
    except ValueError:
        return None


def is_pricing_label(value: Any) -> bool:
    label = normalized_text(value)
    return "valor sugerido" in label or "preco sugerido" in label


def suggested_price(rows: list[tuple[Any, ...]]) -> float | None:
    """Return the preferred suggested price from a pricing worksheet."""

    return suggested_price_details(rows)[0]


def suggested_price_details(rows: list[tuple[Any, ...]]) -> tuple[float | None, bool]:
    """Find a price without crossing into an adjacent worksheet table.

    Some source sheets contain an auxiliary delivery table far to the right of
    the pricing area. Only cells adjacent to the label are eligible. A price
    explicitly marked as including fees takes priority over the simple price.
    """

    candidates: list[tuple[int, int, float]] = []
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if not is_pricing_label(value):
                continue
            label = normalized_text(value)
            priority = 2 if "taxa" in label else 1
            for offset in (1, 2):
                candidate_index = column_index + offset
                if candidate_index >= len(row):
                    continue
                numeric = as_number(row[candidate_index])
                if numeric is not None and numeric > 0:
                    candidates.append((priority, offset, numeric))
            for offset in (1, 2):
                next_row_index = row_index + offset
                if next_row_index >= len(rows) or column_index >= len(rows[next_row_index]):
                    continue
                numeric = as_number(rows[next_row_index][column_index])
                if numeric is not None and numeric > 0:
                    candidates.append((priority, offset, numeric))
    if not candidates:
        return None, False
    priority, _, price = sorted(candidates, key=lambda item: (-item[0], item[1], item[2]))[0]
    return price, priority == 2


def product_name(rows: list[tuple[Any, ...]], sheet_name: str) -> str:
    for row in rows[:5]:
        if row and isinstance(row[0], str) and row[0].strip():
            candidate = row[0].strip()
            if normalized_text(candidate) not in {"valor do produto", "ingredientes"}:
                return candidate
    clean_sheet_name = sheet_name.strip()
    parts = clean_sheet_name.split(maxsplit=1)
    if len(parts) == 2 and normalized_text(parts[0]) == "precificacao":
        clean_sheet_name = parts[1]
    return clean_sheet_name or "Produto sem nome"


def product_category(name: str) -> str:
    label = normalized_text(name)
    if "sorvete" in label:
        return "Sorvete"
    if "acai" in label:
        return "Acai"
    return "Produto"


def normalize_pricing_workbook(
    path: Path,
    source_hash: str,
    source_path: str,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    """Normalize product sheets into ingredient and product JSONL records.

    The expected ingredient layout is compatible with the existing pricing
    workbook: description, unit, purchase quantity/value and recipe
    quantity/value. Rows that do not meet this shape are ignored and reported
    at worksheet level instead of causing a partial run to fail.
    """

    workbook = load_workbook(path, data_only=True, read_only=True)
    records: list[dict[str, Any]] = []
    diagnostics: list[dict[str, str]] = []
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                diagnostics.append({"sheet": worksheet.title, "reason": "empty_sheet"})
                continue
            records_for_sheet, sheet_diagnostics = normalize_pricing_sheet(
                rows=rows,
                source_hash=source_hash,
                source_path=source_path,
                sheet_name=worksheet.title,
            )
            records.extend(records_for_sheet)
            diagnostics.extend(sheet_diagnostics)
    finally:
        workbook.close()
    return records, diagnostics


def normalize_pricing_sheet(
    rows: list[tuple[Any, ...]],
    source_hash: str,
    source_path: str,
    sheet_name: str,
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    name = product_name(rows, sheet_name)
    price, price_includes_fees = suggested_price_details(rows)
    ingredients: list[dict[str, Any]] = []
    diagnostics: list[dict[str, str]] = []

    for row_number, row in enumerate(rows, start=1):
        if len(row) < 6 or not isinstance(row[0], str) or not row[0].strip():
            continue
        purchased_quantity = as_number(row[2])
        purchase_value = as_number(row[3])
        used_quantity = as_number(row[4])
        reported_usage_cost = as_number(row[5])
        if None in (purchased_quantity, purchase_value, used_quantity):
            continue
        if purchased_quantity <= 0 or used_quantity <= 0:
            continue
        ingredient_name = row[0].strip()
        unit_cost = purchase_value / purchased_quantity
        usage_cost = unit_cost * used_quantity
        ingredient = {
            "record_type": "pricing_ingredient",
            "transformation_version": TRANSFORMATION_VERSION,
            "source_hash": source_hash,
            "source_path": source_path,
            "source_sheet": sheet_name,
            "source_row": row_number,
            "product_id": f"product-{slug(name)}",
            "product_name": name,
            "ingredient_id": f"ingredient-{slug(ingredient_name)}",
            "ingredient_name": ingredient_name,
            "unit": str(row[1] or "un").strip() or "un",
            "purchase_quantity": round(purchased_quantity, 6),
            "purchase_value": round(purchase_value, 6),
            "used_quantity": round(used_quantity, 6),
            "unit_cost": round(unit_cost, 6),
            "usage_cost": round(usage_cost, 6),
            "reported_usage_cost": round(reported_usage_cost, 6) if reported_usage_cost is not None else None,
        }
        ingredients.append(ingredient)

    if not ingredients:
        diagnostics.append({"sheet": sheet_name, "reason": "no_recipe_rows"})
        return [], diagnostics
    if price is None:
        diagnostics.append({"sheet": sheet_name, "reason": "suggested_price_not_found"})
    elif not price_includes_fees:
        diagnostics.append({"sheet": sheet_name, "reason": "suggested_price_with_fees_not_found"})

    recipe: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "ingredient_id": "",
            "ingredient_name": "",
            "unit": "un",
            "quantity": 0.0,
            "cost": 0.0,
        }
    )
    for ingredient in ingredients:
        item = recipe[ingredient["ingredient_id"]]
        item["ingredient_id"] = ingredient["ingredient_id"]
        item["ingredient_name"] = ingredient["ingredient_name"]
        item["unit"] = ingredient["unit"]
        item["quantity"] += ingredient["used_quantity"]
        item["cost"] += ingredient["usage_cost"]
    normalized_recipe = [
        {
            **item,
            "quantity": round(item["quantity"], 6),
            "cost": round(item["cost"], 6),
        }
        for _, item in sorted(recipe.items())
    ]
    recipe_cost = round(sum(item["cost"] for item in normalized_recipe), 6)
    product = {
        "record_type": "pricing_product",
        "transformation_version": TRANSFORMATION_VERSION,
        "source_hash": source_hash,
        "source_path": source_path,
        "source_sheet": sheet_name,
        "product_id": f"product-{slug(name)}",
        "product_name": name,
        "category": product_category(name),
        "suggested_price": round(price, 2) if price is not None else None,
        "suggested_price_includes_fees": price_includes_fees,
        "recipe_cost": recipe_cost,
        "gross_margin_percent": round(((price - recipe_cost) / price) * 100, 2) if price else None,
        "ingredient_count": len(normalized_recipe),
        "recipe": normalized_recipe,
    }
    return [*ingredients, product], diagnostics
