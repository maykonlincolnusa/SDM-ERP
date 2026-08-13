import json
from datetime import datetime

from openpyxl import Workbook

from etl import ETLPipeline, ETLPaths


def create_pricing_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Precificacao Acai"
    sheet.append(["Acai teste 500 ml"])
    sheet.append(["Base de acai", "kg", 10, 100, 0.4, 4])
    sheet.append(["Granola", "kg", 2, 30, 0.05, 0.75])
    sheet.append([None, None, None, None, None, None, None, "VALOR SUGERIDO C/ TAXAS"])
    sheet.append([None, None, None, None, None, None, None, 18])
    workbook.save(path)


def create_financial_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Julho 26"
    sheet.append(["DATA", "DESCRICAO", "VALOR", None, None, None, "DATA", "DINHEIRO", "CARTAO", "PIX", "99 IFOOD", "TOTAL"])
    sheet.append([datetime(2026, 7, 1), "Compra de insumos", 15, None, None, None, 1, 10, 20, 20, 10, 60])
    personal = workbook.create_sheet("CONTROLE PESSOAL")
    personal.append(["DATA", "DESCRICAO", "VALOR"])
    personal.append([datetime(2026, 7, 1), "Despesa pessoal", 999])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def create_market_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pagina1"
    sheet.append(["Localidade", "Produto", "Base", "Frutas", "Padrao social", "Colocacao"])
    sheet.append(["Centro", "Acai", "Colibr roxo", "maca; banana", "Alta", 1])
    sheet.append([None, None, None, None, None, 0])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def test_pricing_workbook_creates_bronze_silver_gold_and_manifest(tmp_path):
    paths = ETLPaths.from_data_dir(tmp_path / "data")
    paths.raw.mkdir(parents=True)
    create_pricing_workbook(paths.raw / "pricing.xlsx")

    result = ETLPipeline(paths).run()

    assert result.processed_files == 1
    assert result.failed_files == 0
    assert result.pricing_records == 3
    assert result.gold_products == 1
    manifest = json.loads((paths.state / "manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["files"]) == 1
    source_hash, entry = next(iter(manifest["files"].items()))
    assert entry["source_paths"] == ["pricing.xlsx"]
    assert (paths.bronze / "files" / f"{source_hash}.xlsx").exists()
    assert (paths.bronze / "xlsx" / f"{source_hash}.jsonl").exists()
    silver_rows = [json.loads(line) for line in (paths.silver / "pricing" / f"{source_hash}.jsonl").read_text(encoding="utf-8").splitlines()]
    product = next(row for row in silver_rows if row["record_type"] == "pricing_product")
    assert product["suggested_price"] == 18.0
    assert product["recipe_cost"] == 4.75
    assert product["gross_margin_percent"] == 73.61
    catalog = [json.loads(line) for line in (paths.gold / "pricing_catalog.jsonl").read_text(encoding="utf-8").splitlines()]
    assert catalog == [
        {
            "category": "Acai",
            "gross_margin_percent": 73.61,
            "ingredient_count": 2,
            "product_id": "product-acai-teste-500-ml",
            "product_name": "Acai teste 500 ml",
            "recipe": [
                {"cost": 4.0, "ingredient_id": "ingredient-base-de-acai", "ingredient_name": "Base de acai", "quantity": 0.4, "unit": "kg"},
                {"cost": 0.75, "ingredient_id": "ingredient-granola", "ingredient_name": "Granola", "quantity": 0.05, "unit": "kg"},
            ],
            "recipe_cost": 4.75,
            "source_hash": source_hash,
            "source_path": "pricing.xlsx",
            "source_sheet": "Precificacao Acai",
            "suggested_price": 18.0,
            "suggested_price_includes_fees": True,
        }
    ]


def test_etl_is_idempotent_for_the_same_workbook(tmp_path):
    paths = ETLPaths.from_data_dir(tmp_path / "data")
    paths.raw.mkdir(parents=True)
    create_pricing_workbook(paths.raw / "pricing.xlsx")

    first = ETLPipeline(paths).run()
    manifest = json.loads((paths.state / "manifest.json").read_text(encoding="utf-8"))
    source_hash = next(iter(manifest["files"]))
    silver_path = paths.silver / "pricing" / f"{source_hash}.jsonl"
    first_content = silver_path.read_text(encoding="utf-8")

    second = ETLPipeline(paths).run()

    assert first.processed_files == 1
    assert second.processed_files == 0
    assert second.skipped_files == 1
    assert second.pricing_records == 3
    assert silver_path.read_text(encoding="utf-8") == first_content
    assert len(list((paths.state / "runs").glob("*.json"))) == 2


def test_finance_and_market_use_separate_silver_contracts(tmp_path):
    paths = ETLPaths.from_data_dir(tmp_path / "data")
    create_financial_workbook(paths.raw / "finance" / "finance.xlsx")
    create_market_workbook(paths.raw / "market_research" / "market.xlsx")

    result = ETLPipeline(paths).run()

    assert result.failed_files == 0
    assert result.financial_records == 5
    assert result.market_records == 1
    assert result.document_records == 2
    daily_revenue = [json.loads(line) for line in (paths.gold / "daily_revenue.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(daily_revenue) == 1
    assert daily_revenue[0]["business_date"] == "2026-07-01"
    assert daily_revenue[0]["total_revenue"] == 60.0
    assert daily_revenue[0]["source_record_count"] == 4
    assert {entry["channel"]: entry["amount"] for entry in daily_revenue[0]["payment_channels"]} == {
        "Cartão": 20.0,
        "Dinheiro": 10.0,
        "Pix": 20.0,
        "iFood": 10.0,
    }
    expenses = [json.loads(line) for line in (paths.gold / "expenses.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(expenses) == 1
    assert expenses[0]["description"] == "Compra de insumos"
    market_summary = json.loads((paths.gold / "market_summary.json").read_text(encoding="utf-8"))
    assert market_summary["observation_count"] == 1
    quality = json.loads((paths.gold / "data_quality.json").read_text(encoding="utf-8"))
    assert quality["issue_count"] >= 3
